"""
sample/sample.xlsx 서식/메타데이터 분석 스크립트 (일회성, 데이터 값 자체는 분석 대상 아님).

색상(배경색/폰트색)은 의도적으로 출력에서 제외한다 (사용자 요청).
분석 항목: 시트 구성, 컬럼(헤더), 폰트, 테두리, 정렬/줄바꿈, 열 너비,
틀 고정, 오토필터, 조건부 서식 규칙(색상 제외), 데이터 검증(드롭다운).
"""

import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SAMPLE_PATH = PROJECT_ROOT / "sample" / "sample.xlsx"


def describe_font(font) -> str:
    parts = []
    if font.name:
        parts.append(f"name={font.name}")
    if font.size:
        parts.append(f"size={font.size}")
    if font.bold:
        parts.append("bold")
    if font.italic:
        parts.append("italic")
    if font.underline and font.underline != "none":
        parts.append(f"underline={font.underline}")
    return ", ".join(parts) if parts else "(default)"


def describe_border(border) -> str:
    sides = []
    for side_name in ("left", "right", "top", "bottom"):
        side = getattr(border, side_name)
        if side and side.style:
            sides.append(f"{side_name}={side.style}")
    return ", ".join(sides) if sides else "(none)"


def describe_alignment(align) -> str:
    parts = []
    if align.horizontal:
        parts.append(f"h={align.horizontal}")
    if align.vertical:
        parts.append(f"v={align.vertical}")
    if align.wrap_text:
        parts.append("wrap_text=True")
    return ", ".join(parts) if parts else "(default)"


def analyze_sheet(ws):
    print(f"\n{'=' * 70}")
    print(f"시트: {ws.title}")
    print(f"{'=' * 70}")
    print(f"차원: {ws.dimensions} (max_row={ws.max_row}, max_col={ws.max_column})")
    print(f"틀 고정(freeze_panes): {ws.freeze_panes}")
    print(f"오토필터: {ws.auto_filter.ref}")
    print(f"시트 뷰 - showGridLines: {ws.sheet_view.showGridLines}")

    # 헤더 행 (1행) 분석
    print("\n--- 헤더(1행) ---")
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        col_letter = get_column_letter(col_idx)
        width = ws.column_dimensions[col_letter].width if col_letter in ws.column_dimensions else None
        print(
            f"  {col_letter}1: value={cell.value!r} | font=[{describe_font(cell.font)}] "
            f"| border=[{describe_border(cell.border)}] | align=[{describe_alignment(cell.alignment)}] "
            f"| number_format={cell.number_format!r} | col_width={width}"
        )

    # 행 높이 샘플
    print("\n--- 행 높이(처음 5행) ---")
    for r in range(1, min(6, ws.max_row + 1)):
        rd = ws.row_dimensions.get(r)
        print(f"  row {r}: height={rd.height if rd else None}")

    # 데이터 행 스타일 샘플 (2~4행, 모든 컬럼)
    print("\n--- 데이터 샘플 (2~4행) ---")
    for r in range(2, min(5, ws.max_row + 1)):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=col_idx)
            col_letter = get_column_letter(col_idx)
            print(
                f"  {col_letter}{r}: value={cell.value!r} | font=[{describe_font(cell.font)}] "
                f"| border=[{describe_border(cell.border)}] | align=[{describe_alignment(cell.alignment)}] "
                f"| number_format={cell.number_format!r}"
            )
        print("  ---")

    # 병합 셀
    print(f"\n--- 병합 셀 ---\n  {list(ws.merged_cells.ranges)}")

    # 조건부 서식 (규칙 종류만, 색상 제외)
    print("\n--- 조건부 서식 규칙 ---")
    cf = ws.conditional_formatting
    if not list(cf):
        print("  (없음)")
    else:
        for rule_holder in cf:
            print(f"  범위: {rule_holder.sqref}")
            for rule in rule_holder.rules:
                info = {
                    "type": rule.type,
                    "operator": getattr(rule, "operator", None),
                    "formula": getattr(rule, "formula", None),
                    "priority": rule.priority,
                    "stopIfTrue": rule.stopIfTrue,
                }
                print(f"    - {info}")

    # 데이터 검증 (드롭다운 등)
    print("\n--- 데이터 검증(Data Validation) ---")
    dvs = list(ws.data_validations.dataValidation)
    if not dvs:
        print("  (없음)")
    else:
        for dv in dvs:
            print(
                f"  type={dv.type}, formula1={dv.formula1!r}, "
                f"allow_blank={dv.allow_blank}, ranges={dv.sqref}"
            )

    # 열 너비 전체
    print("\n--- 열 너비 전체 ---")
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        cd = ws.column_dimensions.get(col_letter)
        print(f"  {col_letter}: width={cd.width if cd else None}")


def main():
    if not SAMPLE_PATH.exists():
        print(f"샘플 파일을 찾을 수 없습니다: {SAMPLE_PATH}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(str(SAMPLE_PATH), data_only=False)
    print(f"파일: {SAMPLE_PATH}")
    print(f"시트 목록: {wb.sheetnames}")

    for ws in wb.worksheets:
        analyze_sheet(ws)


if __name__ == "__main__":
    main()
