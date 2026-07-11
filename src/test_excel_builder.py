"""
excel_builder.py Dry-run 테스트.

더미 문제/검증 데이터로 실제 xlsx 파일을 생성하고, openpyxl로 다시 열어
시트 구성/행 수/헤더/데이터 유효성을 검증한다.
"""

import logging
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))

import openpyxl

from excel_builder import ExcelBuilder

OUTPUT_DIR = PROJECT_ROOT / "output" / "_dryrun_test"


def _logger() -> logging.Logger:
    logger = logging.getLogger("excel_dryrun")
    if not logger.handlers:
        handler = logging.StreamHandler(sys.stdout)
        handler.setFormatter(logging.Formatter("[%(levelname)s] %(name)s: %(message)s"))
        logger.addHandler(handler)
    logger.setLevel(logging.INFO)
    return logger


def _dummy_questions():
    return [
        {
            "id": 1,
            "text": "다음 중 계약의 성립 요건이 아닌 것은?",
            "options": [
                {"num": "①", "text": "청약", "correct": False, "explanation": "해설 없음"},
                {"num": "②", "text": "승낙", "correct": False, "explanation": "해설 없음"},
                {"num": "③", "text": "물권 설정", "correct": True, "explanation": "정답: 물권 설정은 계약 성립요건이 아님"},
                {"num": "④", "text": "합의", "correct": False, "explanation": "해설 없음"},
                {"num": "⑤", "text": "당사자 확정", "correct": False, "explanation": "해설 없음"},
            ],
            "confidence": 1.0,
            "status": "AUTO",
            "chunk_id": "chunk_001",
        },
        {
            "id": 2,
            "text": "다음 중 물권이 아닌 것은?",
            "options": [
                {"num": "①", "text": "소유권", "correct": False, "explanation": "해설 없음"},
                {"num": "②", "text": "저당권", "correct": False, "explanation": "해설 없음"},
                {"num": "③", "text": "채권", "correct": True, "explanation": "정답: 채권은 물권이 아님"},
            ],
            "confidence": 0.6,
            "status": "HUMAN_REVIEW",
            "chunk_id": "chunk_001",
        },
        {
            "id": 3,
            "text": "다음 중 옳은 것은?",
            "options": [
                {"num": "①", "text": "가", "correct": False, "explanation": "해설 없음"},
                {"num": "②", "text": "나", "correct": True, "explanation": "정답"},
            ],
            "confidence": 0.9,
            "status": "HUMAN_REVIEW",
            "chunk_id": "chunk_002",
        },
    ]


def test_create_all_sheets():
    """전체 시트 생성 후 파일이 실제로 열리고 구조가 기대와 일치하는지 확인."""
    builder = ExcelBuilder(str(OUTPUT_DIR), logger=_logger())
    all_questions = _dummy_questions()
    validation_results = {
        "human_review": [q for q in all_questions if q["status"] == "HUMAN_REVIEW"],
        "statistics": {
            "total_problems": 3,
            "auto_confirmed_ratio": 1 / 3,
            "human_review_ratio": 2 / 3,
            "error_breakdown": {"OPTION_COUNT_INCONSISTENT": 2, "TYPO_REGRESSION": 1},
            "total_duration_seconds": 42.5,
            "failed_chunks": 0,
        },
    }

    output_files = builder.create_all_sheets(all_questions, validation_results)
    assert "main" in output_files
    path = output_files["main"]
    assert path.exists(), f"파일이 생성되지 않음: {path}"
    print(f"  PASS: 파일 생성 확인 ({path})")

    wb = openpyxl.load_workbook(str(path))
    assert wb.sheetnames == ["전체 문제", "검토 필요", "변환 통계"], wb.sheetnames
    print(f"  PASS: 시트 구성 확인 {wb.sheetnames}")

    ws_main = wb["전체 문제"]
    # 헤더: A 번호 | B NO. | C 문제 | D 선지번호 | E 선지 내용 | F (spacer) | G 정오(O/X) | H 해설 | I 신뢰도 | J 상태
    assert ws_main["A1"].value == "번호"
    assert ws_main["B1"].value == "NO."
    assert ws_main["C1"].value == "문제"
    assert not ws_main["F1"].value  # spacer 헤더는 공란(빈 문자열이 저장 시 None으로 정규화됨)
    assert ws_main["G1"].value == "정오(O/X)"
    assert ws_main["I1"].value == "신뢰도"
    assert ws_main["J1"].value == "상태"
    total_option_rows = sum(len(q["options"]) for q in all_questions)
    assert ws_main.max_row == 1 + total_option_rows, (
        f"전체 문제 시트 행 수 불일치: {ws_main.max_row} != {1 + total_option_rows}"
    )
    assert ws_main.freeze_panes == "A2"
    assert ws_main.auto_filter.ref is not None
    print(f"  PASS: 전체 문제 시트 - 헤더(A~J)/{total_option_rows}개 선지 행/틀고정/필터 확인")

    assert not ws_main["A2"].value  # 번호 컬럼은 수동 작업용 공란(빈 문자열→저장 시 None)
    assert ws_main["B2"].value == 1   # NO.(문제번호)
    assert ws_main["C2"].value == "다음 중 계약의 성립 요건이 아닌 것은?"
    assert ws_main["C3"].value == "동일"  # 2번째 선지부터 "동일"
    assert ws_main["F2"].value is None    # spacer는 항상 공란
    assert ws_main["G4"].value == "O"     # 정답 선지는 O
    assert ws_main["G2"].value == "X"
    assert ws_main["I2"].value == 1.0     # 신뢰도
    assert ws_main["J2"].value == "AUTO"  # 상태
    print("  PASS: 번호(공란)/NO./'동일' 처리/spacer/정오(O/X)/신뢰도/상태 값 확인")

    ws_review = wb["검토 필요"]
    assert ws_review["A1"].value == "번호"
    assert ws_review["K1"].value == "수동 확인"
    review_option_rows = sum(
        len(q["options"]) for q in all_questions if q["status"] == "HUMAN_REVIEW"
    )
    assert ws_review.max_row == 1 + review_option_rows
    assert ws_review["K2"].value == "미확인"
    dv_list = list(ws_review.data_validations.dataValidation)
    assert len(dv_list) == 1
    assert "O,X" in dv_list[0].formula1
    print(f"  PASS: 검토 필요 시트 - {review_option_rows}개 행, 드롭다운 검증 컬럼(K) 확인")

    # 서식: 폰트(마루 부리 중간)/전 셀 테두리(thin) 적용 확인 (색상 값은 검증 대상 아님)
    header_cell = ws_main["C1"]
    assert header_cell.font.name == "마루 부리 중간"
    assert header_cell.font.bold is True
    assert header_cell.border.left.style == "thin"
    data_cell = ws_main["C2"]
    assert data_cell.font.name == "마루 부리 중간"
    assert data_cell.border.top.style == "thin"
    assert data_cell.alignment.vertical == "center"
    assert data_cell.alignment.wrap_text is True
    short_cell = ws_main["B2"]
    assert short_cell.alignment.horizontal == "center"
    print("  PASS: 폰트(마루 부리 중간)/테두리(thin)/정렬(wrap+center) 서식 확인")

    ws_stats = wb["변환 통계"]
    stats_values = [row[0].value for row in ws_stats.iter_rows(min_row=2, max_col=1)]
    assert "전체 문제 수" in stats_values
    assert "오류 유형" in stats_values
    print(f"  PASS: 변환 통계 시트 - {stats_values}")

    wb.close()


def test_empty_review():
    """검토 필요 항목이 0건이어도 오류 없이 빈 시트가 생성되어야 함."""
    builder = ExcelBuilder(str(OUTPUT_DIR), logger=_logger())
    all_questions = [q for q in _dummy_questions() if q["status"] == "AUTO"]
    validation_results = {
        "human_review": [],
        "statistics": {
            "total_problems": 1, "auto_confirmed_ratio": 1.0, "human_review_ratio": 0.0,
            "error_breakdown": {}, "total_duration_seconds": 5.0, "failed_chunks": 0,
        },
    }
    output_files = builder.create_all_sheets(all_questions, validation_results)
    wb = openpyxl.load_workbook(str(output_files["main"]))
    ws_review = wb["검토 필요"]
    assert ws_review.max_row == 1  # 헤더만
    print("  PASS: 검토 필요 0건 -> 헤더만 있는 시트 생성 확인")
    wb.close()


def main():
    tests = [
        ("create_all_sheets (기본 3건)", test_create_all_sheets),
        ("create_all_sheets (검토 필요 0건)", test_empty_review),
    ]
    failed = []
    for name, fn in tests:
        print(f"\n=== {name} ===")
        try:
            fn()
        except AssertionError as e:
            failed.append(name)
            print(f"  FAIL: {e}")

    print("\n" + "=" * 50)
    if failed:
        print(f"실패: {len(failed)}/{len(tests)} - {failed}")
        sys.exit(1)
    print(f"전체 통과: {len(tests)}/{len(tests)}")


if __name__ == "__main__":
    main()
